import io
import csv
from datetime import datetime, timezone
from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for
from flask_login import login_required, current_user
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import db, Case, SheetConfig, User

reports_bp = Blueprint("reports", __name__)


def build_filtered_query(args):
    """Construye la query de casos filtrada según los parámetros recibidos."""
    query = Case.query

    fecha_desde = args.get("fecha_desde", "").strip()
    fecha_hasta = args.get("fecha_hasta", "").strip()
    sheet_id = args.get("sheet_id", type=int)
    status = args.get("status", "").strip()
    assigned_to = args.get("assigned_to", type=int)
    created_by = args.get("created_by", type=int)
    only_rejected = args.get("only_rejected") == "1"
    only_audited = args.get("only_audited") == "1"

    if fecha_desde:
        try:
            dt_desde = datetime.strptime(fecha_desde, "%Y-%m-%d")
            query = query.filter(Case.created_at >= dt_desde)
        except ValueError:
            pass

    if fecha_hasta:
        try:
            # Fin del día
            dt_hasta = datetime.strptime(fecha_hasta, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
            query = query.filter(Case.created_at <= dt_hasta)
        except ValueError:
            pass

    if sheet_id:
        query = query.filter(Case.sheet_config_id == sheet_id)

    if status:
        query = query.filter(Case.status == status)

    if assigned_to:
        query = query.filter(Case.assigned_to == assigned_to)

    if created_by:
        query = query.filter(Case.created_by == created_by)

    if only_rejected:
        query = query.filter(Case.status == "rechazado")

    if only_audited:
        query = query.filter(Case.quality_score.isnot(None))

    return query


@reports_bp.route("/reportes")
@login_required
def index():
    """Pantalla de configuración y previsualización de reportes."""
    if not current_user.can_export_reports:
        flash("No tenés permisos para acceder a informes.", "error")
        return redirect(url_for("dashboard.index"))

    sheets = SheetConfig.query.filter_by(is_active=True).order_by(SheetConfig.display_name).all()
    back_agents = User.query.filter(User.role.in_(["agente_back", "agent"])).order_by(User.display_name).all()
    front_agents = User.query.filter(User.role.in_(["agente_front"])).order_by(User.display_name).all()

    query = build_filtered_query(request.args)
    total_count = query.count()
    cases_preview = query.order_by(Case.created_at.desc()).limit(30).all()

    return render_template(
        "reports/index.html",
        sheets=sheets,
        back_agents=back_agents,
        front_agents=front_agents,
        total_count=total_count,
        cases_preview=cases_preview,
        filters=request.args,
    )


@reports_bp.route("/reportes/exportar/excel")
@login_required
def export_excel():
    """Descargar informe en formato Excel (.xlsx)."""
    if not current_user.can_export_reports:
        flash("No tenés permisos para exportar informes.", "error")
        return redirect(url_for("dashboard.index"))

    query = build_filtered_query(request.args)
    cases = query.order_by(Case.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Casos CRM BGH"

    # Estilos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = [
        "ID Caso",
        "Tipo de Caso",
        "Estado",
        "N° Pedido",
        "Fecha Creación",
        "Tienda",
        "Solicitud",
        "DNI / CUIT",
        "Tracking",
        "Cliente",
        "Email",
        "Agente Front",
        "Agente Back",
        "Código SAP",
        "Caso Salesforce",
        "Motivo Rechazo",
        "Motivo Re-verificación",
        "Respuesta Front",
        "Calificación TYQ (%)",
        "Feedback Calidad",
        "Auditado Por",
        "Fecha Auditoría",
    ]

    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    for c in cases:
        row = [
            f"#{c.id}",
            c.sheet_config.display_name if c.sheet_config else "",
            c.status_label,
            c.pedido_id or "",
            c.created_at.strftime("%d/%m/%Y %H:%M") if c.created_at else "",
            c.tienda or "",
            c.solicitud or "",
            c.dni_cuit or "",
            c.tracking or "",
            c.nombre_cliente or "",
            c.email_cliente or "",
            c.agente_front or (c.creator.display_name if c.creator else ""),
            c.assigned_user.display_name if c.assigned_user else "Sin Asignar",
            c.codigo_sap or "",
            c.caso_salesforce or "",
            c.motivo_rechazo or "",
            c.motivo_reverificacion or "",
            c.respuesta_front or "",
            f"{c.quality_score}%" if c.quality_score is not None else "",
            c.quality_feedback or "",
            c.auditor.display_name if c.auditor else "",
            c.audited_at.strftime("%d/%m/%Y %H:%M") if c.audited_at else "",
        ]
        ws.append(row)

    # Auto-ajuste de anchos de columna
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell.font = data_font
            cell.border = thin_border
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # Freeze header row
    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Informe_CRM_BGH_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@reports_bp.route("/reportes/exportar/csv")
@login_required
def export_csv():
    """Descargar informe en formato CSV con UTF-8 BOM para apertura directa en Excel."""
    if not current_user.can_export_reports:
        flash("No tenés permisos para exportar informes.", "error")
        return redirect(url_for("dashboard.index"))

    query = build_filtered_query(request.args)
    cases = query.order_by(Case.created_at.desc()).all()

    output = io.StringIO()
    # UTF-8 BOM
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";")

    headers = [
        "ID Caso", "Tipo de Caso", "Estado", "N° Pedido", "Fecha", "Tienda",
        "Solicitud", "DNI/CUIT", "Tracking", "Cliente", "Email", "Agente Front",
        "Agente Back", "Código SAP", "Caso Salesforce", "Motivo Rechazo",
        "Motivo Re-verificación", "Respuesta Front", "Score TYQ", "Feedback Calidad"
    ]
    writer.writerow(headers)

    for c in cases:
        writer.writerow([
            f"#{c.id}",
            c.sheet_config.display_name if c.sheet_config else "",
            c.status_label,
            c.pedido_id or "",
            c.created_at.strftime("%d/%m/%Y %H:%M") if c.created_at else "",
            c.tienda or "",
            c.solicitud or "",
            c.dni_cuit or "",
            c.tracking or "",
            c.nombre_cliente or "",
            c.email_cliente or "",
            c.agente_front or (c.creator.display_name if c.creator else ""),
            c.assigned_user.display_name if c.assigned_user else "",
            c.codigo_sap or "",
            c.caso_salesforce or "",
            c.motivo_rechazo or "",
            c.motivo_reverificacion or "",
            c.respuesta_front or "",
            f"{c.quality_score}%" if c.quality_score is not None else "",
            c.quality_feedback or "",
        ])

    mem = io.BytesIO()
    mem.write(output.getvalue().encode("utf-8-sig"))
    mem.seek(0)

    filename = f"Informe_CRM_BGH_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return send_file(
        mem,
        as_attachment=True,
        download_name=filename,
        mimetype="text/csv",
    )
